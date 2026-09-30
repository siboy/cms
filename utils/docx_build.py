"""
Builder: blok terstruktur (docx_blocks) -> DOCX baru yang rapi.

Yang dirapikan saat build (bukan menyalin format asli):
  - Style seragam (font, heading H1..H4, caption, tabel), A4, margin baku.
  - Penomoran heading konsisten (BAB I., 1.1, 1.1.1) dari struktur, bukan ketikan manual.
  - Paging: H1 mulai halaman baru, judul tidak terpisah dari isi (keep-with-next),
    baris tabel tidak terpotong, header tabel berulang, gambar+caption satu halaman.
  - Seksi: cover (tanpa nomor) -> pendahuluan (i, ii, ...) -> isi (1, 2, ...) + footer dokumen.
  - Daftar Isi/Tabel/Gambar = field Word (TOC/SEQ) sehingga nomor halaman akurat setelah
    dokumen dibuka (Word akan menawarkan "update fields").
  - Caption Tabel/Gambar diberi nomor urut otomatis (SEQ).
"""
from __future__ import annotations

import os
import re
from typing import Optional

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor
from docx.table import _Cell

PAGE_W, PAGE_H = Cm(21.0), Cm(29.7)
MARGIN_L, MARGIN_R, MARGIN_T, MARGIN_B = Cm(3.0), Cm(2.5), Cm(2.5), Cm(2.5)
TEXT_W = PAGE_W - MARGIN_L - MARGIN_R
MAX_IMG_H = Cm(17.0)
ALIGN = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
         "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
ROMAN = [(1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"),
         (50, "L"), (40, "XL"), (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]


def roman(n: int) -> str:
    out = ""
    for v, s in ROMAN:
        while n >= v:
            out += s
            n -= v
    return out


def fmt_num(n: int, nf: str) -> str:
    return {"upperRoman": roman(n), "lowerRoman": roman(n).lower(),
            "upperLetter": chr(64 + n), "lowerLetter": chr(96 + n)}.get(nf, str(n))


# ---------------------------------------------------------------- inline markup
_MARKS = {"**": "b", "__": "i", "++": "u", "^^": "sup", "~~": "sub"}


def parse_inline(s: str):
    """markup -> [(teks, {flags}, url|None)]"""
    out, buf, flags, i = [], [], set(), 0

    def flush(url=None):
        if buf:
            out.append(("".join(buf), set(flags), url))
            buf.clear()

    while i < len(s):
        c = s[i]
        if c == "\\" and i + 1 < len(s):
            buf.append(s[i + 1])
            i += 2
            continue
        if s[i:i + 2] in _MARKS:
            flush()
            f = _MARKS[s[i:i + 2]]
            flags.symmetric_difference_update({f})
            i += 2
            continue
        if c == "[":
            m = re.match(r"\[((?:\\.|[^\]\\])*)\]\(([^)]*)\)", s[i:])
            if m:
                flush()
                for t, f, _ in parse_inline(m.group(1)):
                    out.append((t, f | flags, m.group(2).replace("%29", ")")))
                i += m.end()
                continue
        buf.append(c)
        i += 1
    flush()
    return out


def add_field(par, instr: str, cached: str = "", run_fmt=None):
    def r():
        run = par.add_run()
        if run_fmt:
            run_fmt(run)
        return run
    a = r()
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "begin"); a._r.append(e)
    b = r()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; b._r.append(it)
    c = r()
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "separate"); c._r.append(e)
    d = r(); d.text = cached
    f = r()
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); f._r.append(e)


def add_markup(par, markup: str, size: Optional[Pt] = None, bold: Optional[bool] = None):
    for txt, flags, url in parse_inline(markup):
        parts = re.split(r"(\n|\t)", txt)
        holder = par
        if url:
            rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
            hl = OxmlElement("w:hyperlink"); hl.set(qn("r:id"), rid)
            par._p.append(hl)
        for part in parts:
            if part == "":
                continue
            run = par.add_run()
            if url:
                hl.append(run._r)
                run.font.underline = True
                run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
            if part == "\n":
                run.add_break()
            elif part == "\t":
                run.add_tab()
            else:
                run.text = part
            if "b" in flags or bold:
                run.bold = True
            if "i" in flags:
                run.italic = True
            if "u" in flags and not url:
                run.underline = True
            if "sup" in flags:
                run.font.superscript = True
            if "sub" in flags:
                run.font.subscript = True
            if size:
                run.font.size = size


# ---------------------------------------------------------------- style
def _font(style, name, size=None, bold=None, italic=None, color=RGBColor(0, 0, 0)):
    f = style.font
    f.name = name
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in list(rf.attrib):
        if "Theme" in a or "theme" in a:
            del rf.attrib[a]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color


def setup_styles(doc, base_font: str):
    st = doc.styles
    _font(st["Normal"], base_font, 11)
    pf = st["Normal"].paragraph_format
    pf.space_after, pf.space_before, pf.line_spacing = Pt(6), Pt(0), 1.15
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    spec = {1: (14, True, False, 0, 12), 2: (12.5, True, False, 14, 6),
            3: (11.5, True, False, 10, 4), 4: (11, True, True, 8, 4)}
    for lvl, (sz, b, i, before, after) in spec.items():
        s = st[f"Heading {lvl}"]
        _font(s, base_font, sz, b, i)
        p = s.paragraph_format
        p.space_before, p.space_after, p.keep_with_next, p.keep_together = Pt(before), Pt(after), True, True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if lvl == 1 else WD_ALIGN_PARAGRAPH.LEFT
        p.page_break_before = lvl == 1
        p.line_spacing = 1.0
    cap = st["Caption"]
    _font(cap, base_font, 10.5, True, False)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before, cap.paragraph_format.space_after = Pt(6), Pt(6)
    for n in ("List Bullet", "List Number"):
        st[n].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        st[n].paragraph_format.space_after = Pt(3)


def insert_before(parent, el, successors):
    """Sisipkan el sebelum anak pertama yang tag-nya ada di `successors` (jaga urutan skema OOXML)."""
    for ch in parent:
        if ch.tag in {qn(t) for t in successors}:
            ch.addprevious(el)
            return
    parent.append(el)


PPR_AFTER_BDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                 "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
                 "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                 "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
                 "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
SECT_AFTER_PGNUM = ("w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg", "w:textDirection",
                    "w:bidi", "w:rtlGutter", "w:docGrid", "w:printerSettings", "w:sectPrChange")
SETTINGS_AFTER_UPDATE = ("w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars", "w:rsids",
                         "m:mathPr", "w:attachedSchema", "w:themeFontLang", "w:clrSchemeMapping",
                         "w:doNotIncludeSubdocsInStats", "w:doNotAutoCompressPictures", "w:forceUpgrade",
                         "w:captions", "w:readModeInkLockDown", "w:smartTagType", "w:schemaLibrary",
                         "w:shapeDefaults", "w:doNotEmbedSmartTags", "w:decimalSymbol", "w:listSeparator")


def _shade(cell, fill="D9D9D9"):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _row_flags(row, header: bool, keep: bool):
    trpr = row._tr.get_or_add_trPr()
    if keep:
        trpr.append(OxmlElement("w:cantSplit"))
    if header:
        trpr.append(OxmlElement("w:tblHeader"))


def _new_restart_num(doc, style_name: str):
    """w:num baru untuk daftar bernomor supaya penomoran mulai dari 1 lagi."""
    try:
        numbering = doc.part.numbering_part.element
        base = doc.styles[style_name].element.pPr.numPr.numId.val
        abs_id = numbering.num_having_numId(base).abstractNumId.val
        num = numbering.add_num(abs_id)
        ov = num.add_lvlOverride(ilvl=0)
        ov.add_startOverride(1)
        return num.numId
    except Exception:
        return None


# ---------------------------------------------------------------- label pra-hitung
def compute_labels(blocks: list[dict]) -> dict[int, str]:
    """seq -> label (prefix nomor heading / 'Tabel n' / 'Gambar n')."""
    labels: dict[int, str] = {}
    chap = 0
    sub = [0, 0, 0]
    seqno = {"tabel": 0, "gambar": 0}
    for b in blocks:
        if b["data"].get("skip_build"):
            continue
        if b["kind"] == "heading":
            lvl = b["level"]
            if lvl == 1:
                if b["data"].get("numbered"):
                    chap += 1
                    sub = [0, 0, 0]
                    lab = b["data"].get("num_text") or "%1."
                    labels[b["seq"]] = lab.replace("%1", fmt_num(chap, b["data"].get("num_fmt", "decimal")))
            elif chap and lvl <= 4:
                i = lvl - 2
                sub[i] += 1
                for j in range(i + 1, 3):
                    sub[j] = 0
                labels[b["seq"]] = ".".join([str(chap)] + [str(x) for x in sub[: i + 1]])
        elif b["kind"] == "caption":
            st = b["data"].get("subtype", "tabel")
            seqno[st] += 1
            labels[b["seq"]] = f"{st.capitalize()} {seqno[st]}"
    return labels


class Builder:
    def __init__(self, blocks, assets, meta, media_root=""):
        self.blocks = sorted(blocks, key=lambda b: b["seq"])
        self.assets = {a["sha1"]: a for a in assets}
        self.meta = meta
        self.media_root = media_root
        self.doc = Document()
        self.labels = compute_labels(self.blocks)
        self.seq_counter = {"tabel": 0, "gambar": 0}
        self.fresh_cells: set[int] = set()
        self._num_ctx = None    # (container id, num id) untuk daftar bernomor yang sedang berjalan
        self.stats = {"blocks": 0, "tables": 0, "images": 0, "missing_images": 0}

    # ---- util
    def _para(self, container, style=None):
        if isinstance(container, _Cell) and id(container._tc) in self.fresh_cells:
            self.fresh_cells.discard(id(container._tc))
            p = container.paragraphs[0]
            if style:
                p.style = style
            return p
        return container.add_paragraph(style=style)

    def _asset_path(self, sha1):
        a = self.assets.get(sha1)
        if not a:
            return None
        p = a["path"]
        if not os.path.isfile(p) and self.media_root:
            p = os.path.join(self.media_root, a["filename"])
        return p if os.path.isfile(p) else None

    # ---- render satu daftar blok ke container (Document atau _Cell)
    def render(self, container, blocks, in_cell=False):
        size = Pt(10) if in_cell else None
        for idx, b in enumerate(blocks):
            k, d = b["kind"], b.get("data", {})
            if d.get("skip_build"):
                continue
            nxt = blocks[idx + 1] if idx + 1 < len(blocks) else None
            self.stats["blocks"] += 1
            if k != "list_item":
                self._num_ctx = None if not in_cell else self._num_ctx

            if k == "heading":
                self._heading(container, b)
            elif k == "paragraph":
                p = self._para(container)
                p.alignment = ALIGN.get(d.get("align"), None)
                if in_cell:
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.0
                    if d.get("align") != "center":
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_markup(p, b["text"], size)
            elif k == "list_item":
                self._list_item(container, b, size, in_cell)
            elif k == "caption":
                self._caption(container, b, nxt)
            elif k == "note":
                p = self._para(container)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(4)
                add_markup(p, b["text"], Pt(9))
                for r in p.runs:
                    r.italic = True
            elif k == "image":
                self._image(container, b, nxt, in_cell)
            elif k == "table":
                self._table(container, b)
            elif k == "page_break":
                if (nxt and nxt["kind"] == "heading" and nxt["level"] == 1) or not nxt:
                    continue
                self._para(container).add_run().add_break(WD_BREAK.PAGE)

    def _heading(self, container, b):
        lvl = min(max(b["level"], 1), 4)
        p = container.add_paragraph(style=f"Heading {lvl}")
        lab = self.labels.get(b["seq"])
        text = b["text"]
        if lab:
            p.add_run(lab + " ")
        add_markup(p, text)
        if b["data"].get("generated") and container is self.doc:
            self._generated_list(b["data"]["generated"])

    def _caption(self, container, b, nxt):
        st = b["data"].get("subtype", "tabel")
        self.seq_counter[st] += 1
        p = self._para(container, "Caption")
        p.add_run(st.capitalize() + " ")
        add_field(p, f"SEQ {st.capitalize()} \\* ARABIC", str(self.seq_counter[st]))
        if b["text"]:
            p.add_run(". ")
            add_markup(p, b["text"])
        p.paragraph_format.keep_with_next = bool(nxt and nxt["kind"] in ("table", "image")) or st == "tabel"

    def _list_item(self, container, b, size, in_cell):
        d = b["data"]
        ordered = d.get("ordered")
        style = "List Number" if ordered else "List Bullet"
        p = self._para(container, style)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if in_cell else None
        if ordered:
            key = id(container)
            if not self._num_ctx or self._num_ctx[0] != key or d.get("restart"):
                self._num_ctx = (key, _new_restart_num(self.doc, "List Number"))
            nid = self._num_ctx[1]
            if nid is not None:
                numpr = p._p.get_or_add_pPr().get_or_add_numPr()
                numpr.get_or_add_ilvl().val = 0
                numpr.get_or_add_numId().val = nid
        else:
            self._num_ctx = None
        if d.get("ilvl"):
            p.paragraph_format.left_indent = Cm(0.63 * (d["ilvl"] + 1))
        add_markup(p, b["text"], size)

    def _image(self, container, b, nxt, in_cell):
        d = b["data"]
        path = self._asset_path(d.get("asset"))
        p = self._para(container)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.keep_with_next = bool(nxt and nxt["kind"] in ("caption", "note", "image"))
        if not path:
            self.stats["missing_images"] += 1
            p.add_run(f"[gambar hilang: {d.get('asset', '')[:8]}]")
            return
        a = self.assets[d["asset"]]
        w = Emu(d["cx"]) if d.get("cx") else Cm(a["px_w"] / 96 * 2.54) if a.get("px_w") else TEXT_W
        maxw = Cm(7.5) if in_cell else TEXT_W
        w = min(w, maxw)
        if a.get("px_w") and a.get("px_h") and w * a["px_h"] / a["px_w"] > MAX_IMG_H:
            w = int(MAX_IMG_H * a["px_w"] / a["px_h"])
        p.add_run().add_picture(path, width=w)
        self.stats["images"] += 1

    def _table(self, container, b):
        d = b["data"]
        rows, ncols = d["rows"], max(d["ncols"], 1)
        if not rows:
            return
        self.stats["tables"] += 1
        top = container is self.doc
        t = container.add_table(rows=len(rows), cols=ncols) if top else container.add_table(len(rows), ncols)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        grid = d.get("grid") or []
        if top:
            t.autofit = False
            tot = sum(grid) if len(grid) == ncols and sum(grid) else 0
            widths = [int(TEXT_W * g / tot) for g in grid] if tot else [int(TEXT_W / ncols)] * ncols
        merges = []
        for ri, row in enumerate(rows):
            text_len = sum(len(x["text"]) for c in row["cells"] for x in c["blocks"])
            _row_flags(t.rows[ri], row["header"], keep=text_len < 1500)
            for c in row["cells"]:
                if c["col"] >= ncols:
                    continue
                cell = t.cell(ri, c["col"])
                self.fresh_cells.add(id(cell._tc))
                saved = self._num_ctx
                self.render(cell, c["blocks"], in_cell=True)
                self._num_ctx = saved
                self.fresh_cells.discard(id(cell._tc))
                if row["header"]:
                    _shade(cell)
                    for p in cell.paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for r in p.runs:
                            r.bold = True
                if c["colspan"] > 1 or c["rowspan"] > 1:
                    merges.append((ri, c))
        for ri, c in merges:
            a = t.cell(ri, c["col"])
            z = t.cell(min(ri + c["rowspan"] - 1, len(rows) - 1), min(c["col"] + c["colspan"] - 1, ncols - 1))
            if a._tc is not z._tc:
                a.merge(z)
        if top:
            for row in t.rows:
                seen = set()
                for ci, cell in enumerate(row.cells):
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    span = int(cell._tc.tcPr.gridSpan.val) if cell._tc.tcPr is not None and cell._tc.tcPr.gridSpan is not None else 1
                    cell.width = sum(widths[ci: ci + span])
        # jarak setelah tabel
        if top:
            sp = self.doc.add_paragraph()
            sp.paragraph_format.space_after = Pt(4)

    # ---- daftar otomatis
    def _generated_list(self, kind):
        doc = self.doc
        if kind == "toc":
            entries = []
            for b in self.blocks:
                if b["kind"] == "heading" and b["level"] <= 3 and not b["data"].get("skip_build") \
                        and not b["data"].get("generated"):
                    entries.append((b["level"], (self.labels.get(b["seq"], "") + " " + b["text"]).strip()))
            instr = 'TOC \\o "1-3" \\h \\z \\u'
        else:
            st = "Tabel" if kind == "tof_tabel" else "Gambar"
            entries = []
            n = 0
            for b in self.blocks:
                if b["kind"] == "caption" and b["data"].get("subtype") == st.lower():
                    n += 1
                    entries.append((1, f"{st} {n}. {b['text']}"))
            instr = f'TOC \\h \\z \\c "{st}"'
        if not entries:
            entries = [(1, "(diperbarui otomatis)")]
        paras = []
        for lvl, txt in entries:
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Cm(0.6 * (lvl - 1))
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.tab_stops.add_tab_stop(TEXT_W, WD_TAB_ALIGNMENT.RIGHT, 1)  # titik-titik
            p.add_run(re.sub(r"\\(.)", r"\1", txt))
            paras.append(p)
        first, last = paras[0], paras[-1]

        def fc(kind_, par, at_start):
            r = OxmlElement("w:r")
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind_); r.append(e)
            return r
        # begin + instr + separate di awal paragraf pertama; end di akhir paragraf terakhir
        b_ = fc("begin", first, True)
        i_ = OxmlElement("w:r"); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; i_.append(it)
        s_ = fc("separate", first, True)
        ppr = first._p.pPr
        anchor = ppr if ppr is not None else None
        pos = 1 if anchor is not None else 0
        for k, el in enumerate((b_, i_, s_)):
            first._p.insert(pos + k, el)
        last._p.append(fc("end", last, False))

    # ---- seksi & footer
    def _section_setup(self, section, part):
        section.page_width, section.page_height = PAGE_W, PAGE_H
        section.left_margin, section.right_margin = MARGIN_L, MARGIN_R
        section.top_margin, section.bottom_margin = MARGIN_T, MARGIN_B
        sp = section._sectPr
        for e in sp.findall(qn("w:pgNumType")):
            sp.remove(e)
        f = section.footer
        if part == "cover":
            f.is_linked_to_previous = False
            return
        f.is_linked_to_previous = False
        for p in list(f.paragraphs)[1:]:
            p._p.getparent().remove(p._p)
        p = f.paragraphs[0]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        pg = OxmlElement("w:pgNumType")
        pg.set(qn("w:fmt"), "lowerRoman" if part == "front" else "decimal")
        pg.set(qn("w:start"), "1")
        insert_before(sp, pg, SECT_AFTER_PGNUM)
        small = lambda run: setattr(run.font, "size", Pt(9))
        if part == "front":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_field(p, "PAGE", "i", small)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.tab_stops.add_tab_stop(TEXT_W, WD_TAB_ALIGNMENT.RIGHT)
            ppr = p._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr"); top = OxmlElement("w:top")
            for a, v in (("val", "single"), ("sz", "4"), ("space", "4"), ("color", "808080")):
                top.set(qn("w:" + a), v)
            bd.append(top); insert_before(ppr, bd, PPR_AFTER_BDR)
            r = p.add_run(self.meta.get("footer_text", "")); r.font.size = Pt(8.5)
            p.add_run("\t")
            add_field(p, "PAGE", "1", small)

    def _cover_block(self, b):
        d = b["data"]
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(min(d.get("gap", 0) * 12, 96))
        p.paragraph_format.space_after = Pt(4)
        add_markup(p, b["text"], Pt(d.get("size") or 14), bold=True if b["seq"] <= 8 else None)

    def build(self, out_path: str):
        doc = self.doc
        setup_styles(doc, self.meta.get("base_font") or "Times New Roman")
        cur_part = None
        buf: list[dict] = []
        parts = ["cover", "front", "body", "lampiran"]

        def flush():
            if not buf:
                return
            self.render(doc, buf)
            buf.clear()

        for b in self.blocks:
            part = b["part"]
            if part != cur_part:
                flush()
                sec_part = "body" if part == "lampiran" else part
                if cur_part is None:
                    self._section_setup(doc.sections[0], sec_part)
                elif not (cur_part == "body" and part == "lampiran"):
                    s = doc.add_section(WD_SECTION.NEW_PAGE)
                    self._section_setup(s, sec_part)
                cur_part = part
            if part == "cover":
                if b["kind"] == "paragraph" and not b["data"].get("skip_build"):
                    self._cover_block(b)
                elif b["kind"] == "image" and not b["data"].get("skip_build"):
                    self.render(doc, [b])
            else:
                buf.append(b)
        flush()

        # field diperbarui otomatis saat dibuka
        settings = doc.settings.element
        uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
        insert_before(settings, uf, SETTINGS_AFTER_UPDATE)
        cp = doc.core_properties
        cp.title = self.meta.get("title") or self.meta.get("footer_text", "")[:120]
        cp.author = self.meta.get("author", "")
        os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
        doc.save(out_path)
        return self.stats


def build_docx(blocks, assets, meta, out_path, media_root=""):
    return Builder(blocks, assets, meta, media_root).build(out_path)
