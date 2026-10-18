"""Uji katalog template outline tanpa NAS/DB:  python3 scripts/tpl_test.py   (butuh python-docx + lxml)."""
import io
import json
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from docx import Document  # noqa: E402

from utils import tpl_catalog as tc  # noqa: E402
from utils.tpl_outline import clean_heading, extract_outline  # noqa: E402

OK = 0


def check(cond, msg):
    global OK
    if not cond:
        print("GAGAL:", msg)
        sys.exit(1)
    OK += 1


def mkdocx(items):
    d = Document()
    for kind, text in items:
        if kind.startswith("h"):
            d.add_heading(text, level=int(kind[1]))
        else:
            d.add_paragraph(text)
    b = io.BytesIO()
    d.save(b)
    b.seek(0)
    return b


def t_classify():
    c = tc.classify
    check(c("Rencana Kegiatan Pertambangan Batuan Komoditas Kerikil/Sirtu di Desa X") == "pertambangan", "tambang")
    check(c("Pembangunan SPBU Pertashop Desa Y") == "migas", "spbu")
    check(c("Pembangunan Jalan Lanai-Batang Kundur") == "jalan_transportasi", "jalan")
    check(c("Perkebunan Kelapa Sawit dan PKS PT Z") == "perkebunan", "sawit")
    check(c("Pembangunan Gudang Beras") == "perdagangan_jasa", "gudang")
    check(c("Kegiatan Budidaya Aneh") == "lainnya", "tak dikenal -> lainnya")
    check(c("Tambang") == "pertambangan" and c("Ketambangan Bertambangan") == "lainnya", "kata utuh, bukan potongan")


def t_clean():
    check(clean_heading("BAB I PENDAHULUAN", 1) == "Pendahuluan", "BAB + huruf besar")
    check(clean_heading("1.2 Latar Belakang", 2) == "Latar Belakang", "nomor manual")
    check(clean_heading("A. Rencana Kegiatan ........ 12", 2) == "Rencana Kegiatan", "titik-titik daftar isi")
    check(clean_heading("PENGELOLAAN LB3 DAN IPAL", 1) == "Pengelolaan LB3 dan IPAL", "akronim & kata sambung")


def t_extract():
    f = mkdocx([("h1", "BAB I PENDAHULUAN"), ("p", "isi"), ("h2", "1.1 Latar Belakang"), ("h2", "1.2 Tujuan"),
                ("h3", "1.2.1 Manfaat"), ("h1", "DAFTAR ISI"), ("h1", "BAB II DESKRIPSI KEGIATAN"), ("h2", "Tabel 2.1 Data"),
                ("h2", "Lokasi"), ("h2", "Jadwal"), ("h2", "Komponen"), ("h1", "Daftar Pustaka"), ("h1", "Lampiran")])
    o = extract_outline(f)
    check(o[0] == (1, "Pendahuluan"), f"H1 dibersihkan: {o[0]}")
    check((3, "Manfaat") in o, "H3 terbaca")
    check(not any("Daftar Isi" in t or "Tabel" in t for _l, t in o), "daftar isi & caption tabel dilewati")
    check(all(o[i][0] <= (o[i - 1][0] + 1 if i else 1) for i in range(len(o))), "tak ada lompat tingkat")
    # dokumen tanpa gaya heading -> pola teks bernomor
    f2 = mkdocx([("p", "BAB I PENDAHULUAN")] + [x for i in range(1, 5) for x in
                 [("p", f"1.{i} Bagian Nomor {i}"), ("p", "isi isi isi")]] + [("p", "BAB II LAIN"), ("p", "2.1 Sub Dua"),
                 ("p", "2.2 Sub Tiga"), ("p", "2.2.1 Anak")])
    o2 = extract_outline(f2)
    check(o2 and o2[0][0] == 1 and any(t == "Sub Dua" for _l, t in o2), f"cadangan pola teks: {o2[:4]}")


def t_consensus():
    def doc(extra=(), drop=()):
        base = [(1, "Pendahuluan"), (2, "Latar Belakang"), (2, "Tujuan"), (1, "Deskripsi Kegiatan"), (2, "Lokasi"),
                (1, "Daftar Pustaka")]
        return [x for x in base if x[1] not in drop] + list(extra)
    outs = [doc(), doc(), doc(drop=("Tujuan",)), doc(extra=[(1, "Rencana Tambang PT Rahasia Jaya")]),
            doc(extra=[(1, "Lampiran")])]
    nodes = tc.consensus(outs)
    ts = [n["t"] for n in nodes]
    check("Rencana Tambang PT Rahasia Jaya" not in ts, "judul spesifik proyek (1 dokumen) tidak bocor")
    check(ts == ["Pendahuluan", "Latar Belakang", "Tujuan", "Deskripsi Kegiatan", "Lokasi", "Daftar Pustaka"], f"urutan: {ts}")
    check(next(n for n in nodes if n["t"] == "Tujuan")["s"] == 80, "persentase dokumen")
    # anak tak boleh hidup tanpa induk
    check(tc.consensus([[(1, "A"), (2, "X")], [(1, "B"), (2, "X")], [(1, "C"), (2, "X")]], min_share=0.5) == [],
          "sub-bagian dgn induk berbeda-beda tak lolos")
    # urutan antar-saudara mengikuti dokumen
    o = [[(1, "Satu"), (1, "Dua"), (1, "Tiga")]] * 3
    check([n["t"] for n in tc.consensus(o)] == ["Satu", "Dua", "Tiga"], "urutan saudara")
    ns = tc.nodes_for({"nodes": [{"l": 1, "t": "A", "s": 30}, {"l": 2, "t": "A1", "s": 90}, {"l": 1, "t": "B", "s": 90},
                                  {"l": 2, "t": "B1", "s": 50}]}, 60)
    check([n["t"] for n in ns] == ["B"], f"ambang kerincian: anak ikut hilang bersama induk: {ns}")


def t_content():
    ap = "PT. Cipta Agro Mitra Mandiri"
    t = tc.templatize("Kegiatan PT CIPTA AGRO MITRA MANDIRI berada di Desa Pinang Damai, Kecamatan Torgamba, Kabupaten Labuhanbatu "
                      "milik PT Sinar Jaya Abadi.", ap)
    check("CIPTA" not in t.upper() and "Sinar" not in t and "Pinang" not in t and "Torgamba" not in t, f"identitas dimasking: {t}")
    check("[Nama Pemrakarsa]" in t and "Desa [nama]" in t and "Kecamatan [nama]" in t, f"placeholder: {t}")

    def doc(name, extra_p=()):
        items = [{"k": "h", "l": 1, "t": "Lokasi"},
                 {"k": "p", "t": f"Kegiatan {name} berada di Desa Aneka dengan luas {len(name)}0 ha dan akan dibangun bertahap."},
                 {"k": "p", "t": "Isi disesuaikan dengan keadaan kegiatan masing-masing."},
                 {"k": "t", "h": [["No", "Jenis", "Luas (ha)"]], "n": 5, "cap": f"Luas {name}"},
                 {"k": "kv", "labels": ["Nama Penanggung Jawab", "Alamat", "Telepon"], "n": 3, "cap": ""},
                 {"k": "h", "l": 1, "t": "Penutup"}]
        return [dict(x) for x in tc.templatize_items(items, name)] + [{"k": "p", "t": p} for p in extra_p]
    docs = [doc(n, ["Kontak: a@b.com untuk pertanyaan lebih lanjut kepada pengelola."]) for n in
            ("PT Alfa Satu", "PT Beta Dua", "PT Gamma Tiga", "PT Delta Empat")]
    nodes = tc.consensus([[(x["l"], x["t"]) for x in d if x["k"] == "h"] for d in docs])
    tc.attach_content(nodes, docs)
    b = nodes[0]["b"]
    kinds = [x["k"] for x in b]
    check(kinds.count("p") == 2 and "kv" in kinds and "t" in kinds, f"jenis blok konsensus: {kinds}")
    check(any(x["k"] == "p" and x["t"].startswith("Kegiatan [Nama Pemrakarsa] berada") for x in b), "teks baku dgn placeholder lolos")
    check(not any("a@b.com" in x.get("t", "") for x in b), "teks berisi email dibuang")
    check(next(x for x in b if x["k"] == "t")["h"] == [["No", "Jenis", "Luas (ha)"]], "header tabel")
    check(next(x for x in b if x["k"] == "kv")["labels"] == ["Nama Penanggung Jawab", "Alamat", "Telepon"], "label tabel isian")
    check("…" in next(x for x in b if x["k"] == "p" and "luas" in x["t"])["t"], "angka beragam diganti …")
    check(tc.nodes_for({"nodes": nodes}, 0)[0].get("b"), "nodes_for membawa blok")


def t_variants():
    a = [(1, "Identitas Penanggung Jawab Usaha dan/atau Kegiatan"), (1, "Matriks Ukl-upl")]
    b = [(1, "Identitas Penanggung Jawab Usaha"), (1, "Matriks UKL-UPL")]
    nodes = tc.consensus([a, b, a, b])
    check(len(nodes) == 2 and nodes[0]["s"] == 100, f"variasi penulisan digabung: {nodes}")
    check(nodes[1]["t"] == "Matriks UKL-UPL", "akronim dirapikan")
    check(tc.fix_display("Batas Waktu (dph) Dan/Atau Amdal") == "Batas Waktu (DPH) dan/atau AMDAL", "fix_display")


def t_search():
    with tempfile.TemporaryDirectory() as td:
        bundle = {"version": 1, "generated": "2026-10-18", "templates": [
            tc.finalize("ukl_upl-pertambangan", "ukl_upl", "pertambangan", [{"l": 1, "t": "Reklamasi Pascatambang", "s": 70}], 12, "konsensus"),
            tc.finalize("andal-pertambangan", "andal", "pertambangan", [{"l": 1, "t": "Pendahuluan", "s": 100}], 9, "konsensus"),
            tc.finalize("ukl_upl-migas", "ukl_upl", "migas", [{"l": 1, "t": "Pendahuluan", "s": 100}], 9, "konsensus"),
            tc.finalize("ka-umum", "ka", "umum", tc.BASELINE_KA, 0, "baku")]}
        p = Path(td) / "t.json"
        p.write_text(json.dumps(bundle), encoding="utf-8")
        tc.BUNDLE, tc._cache["mtime"] = p, None
        r = tc.search("tambang")
        check({x["id"] for x in r} == {"ukl_upl-pertambangan", "andal-pertambangan"}, f"kata kunci 'tambang': {[x['id'] for x in r]}")
        check({x["id"] for x in tc.search("reklamasi")} == {"ukl_upl-pertambangan"}, "cocok di judul bagian")
        check(tc.search("tambang", kind="andal")[0]["id"] == "andal-pertambangan", "filter kind")
        check(tc.search("spbu") and tc.search("spbu")[0]["id"] == "ukl_upl-migas", "sinonim jenis kegiatan (spbu -> migas)")
        check(tc.search("zzzz") == [], "tak ada hasil")
        check(len(tc.search("")) == 4, "tanpa kata kunci = semua")
        check(tc.get("ka-umum")["source"] == "baku" and tc.get("nope") is None, "get")


def t_tools():
    from utils import tpl_tools as tt
    check(tt.match("Rencana Kebutuhan Tenaga Kerja") == ["tenaga_kerja"], "tenaga kerja")
    check("penduduk" in tt.match("Kependudukan dan Proyeksi Jumlah Penduduk"), "penduduk")
    m = tt.match("Pengelolaan Limbah Cair dan IPAL")
    check("air_limbah" in m and "air_sungai" in m, f"limbah cair -> hitung + simulasi: {m}")
    check("bising" in tt.match("Kebisingan dan Getaran") and "udara" in tt.match("Kualitas Udara Ambien"), "bising/udara")
    check(tt.match("Latar Belakang") == [] and tt.match("") == [], "judul umum tidak ditandai")
    try:
        from utils import kd_formulas as kf, kd_sim_schema as ks
        ids = {x.get("id") for x in (kf.CATALOG.values() if isinstance(kf.CATALOG, dict) else kf.CATALOG)} | set(kf.CATALOG if isinstance(kf.CATALOG, dict) else [])
        for t in tt.TOOLS:
            check(all(c in ids for c in t["calc"]), f"id rumus {t['id']} ada di CATALOG")
            check(t.get("tab") in ("proyeksi", "model", "rumus", "sim"), "tab valid")
        sims = {"water", "air", "noise"}
        check(all(t["sim"] in sims for t in tt.TOOLS if t.get("sim")), "id simulator valid")
    except ImportError as e:
        print("lewati uji katalog rumus:", e)
    try:
        from utils import kd_formulas as kf
        d = lambda i: {v["name"]: v["default"] for v in kf.BY_ID[i]["vars"]}
        r = kf.calc("tenaga_kerja_tahap", d("tenaga_kerja_tahap"))
        o = {x["key"]: x["value"] for x in r["outputs"]}
        check(o["lokal"] == 150 and o["pendatang"] == 100 and o["tambah_penduduk"] == 190, f"tenaga kerja tahap: {o}")
        r = kf.calc("tenaga_kerja_tahunan", d("tenaga_kerja_tahunan"))
        check(len(r["table"]["rows"]) == 8 and len(r["series"]["ys"]) == 4 and r["outputs"][0]["value"] == 450, "tenaga kerja tahunan: tabel+deret")
        check(abs(kf.calc("kebutuhan_tk_shift", d("kebutuhan_tk_shift"))["outputs"][0]["value"] - 123.2) < 1e-9, "roster shift")
        check(abs(kf.calc("indeks_pencemaran", d("indeks_pencemaran"))["outputs"][0]["value"] - 1.029) < 0.001, "IP")
        r = kf.calc("lalin_banding", d("lalin_banding")); o = {x["key"]: x["value"] for x in r["outputs"]}
        check(len(r["table"]["rows"]) == 16 and o["ds2"] < o["ds1"] and abs(o["ubah_arus"] - 10) < 1e-9 and r["table"]["rows"][0][7] in "ABCDEF", "lalin sebelum/sesudah")
        o = {x["key"]: x["value"] for x in kf.calc("kepadatan_greenshields", d("kepadatan_greenshields"))["outputs"]}
        check(abs(o["rasio"] - 0.5) < 1e-9 and abs(o["k"] - 17.574) < 0.01, "greenshields")
        try:
            kf.calc("kepadatan_greenshields", {**d("kepadatan_greenshields"), "q": 5000}); check(False, "arus > kapasitas harus ditolak")
        except kf.CalcError:
            check(True, "")
        check("lalu_lintas" in tt.match("Dampak Jembatan terhadap Lalu Lintas") and "lalu_lintas" in tt.match("Jalan Tol"), "penanda jembatan/tol")
        for bad in ("dewatering_dupuit",):
            try:
                kf.calc(bad, {**d(bad), "h": 30}); check(False, "H<=h harus ditolak")
            except kf.CalcError:
                check(True, "")
    except ImportError as e:
        print("lewati uji rumus baru:", e)
    nodes = tc.annotate([{"l": 1, "t": "Kebutuhan Tenaga Kerja", "s": 90}, {"l": 1, "t": "Pendahuluan", "s": 100}], "pertambangan")
    check(nodes[0].get("tools") == ["tenaga_kerja"] and "tools" not in nodes[1], "annotate menambah tools")
    from cmsapp import export
    bl = [{"id": 1, "data": {}}, {"id": 2, "data": {"source": "template_guide"}}, {"id": 3}]
    check([b["id"] for b in export.strip_guide(bl)] == [1, 3], "petunjuk dibuang dari ekspor")


def t_wiring():
    try:
        from cmsapp import templates_api  # noqa: F401
    except ImportError as e:                               # flask/redis tak terpasang di mesin uji
        print("lewati uji wiring:", e)
        return
    check(templates_api.PERM == "template_use", "izin")
    from cmsapp import auth
    check("template_use" in auth.PERM_KEYS, "izin terdaftar di katalog Privilege")


if __name__ == "__main__":
    for fn in (t_classify, t_clean, t_extract, t_consensus, t_content, t_variants, t_search, t_tools, t_wiring):
        fn()
    print(f"OK: {OK} pemeriksaan lulus")
